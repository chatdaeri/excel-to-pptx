#!/usr/bin/env python3
"""
엑셀 스캐너 — deck.cjs 작성용 가벼운 시트 요약

사용:
  python3 scan-sheets.py <엑셀_경로> [--out <저장경로>]

각 시트에서 다음만 추출:
  - 1행 첫 non-empty 셀 → main 메시지
  - 2행 첫 non-empty 셀 → sub 메시지 (없으면 "(없음)")
  - "[표]" 또는 "[차트]" / "[차트:타입]" 으로 시작하는 행 → 블록 제목
    · [표]       → 표로 렌더
    · [차트]      → 차트로 렌더 (default: bar)
    · [차트:line] → 차트 + 타입 지정 (bar | line | doughnut | barLine)
  - 제목 바로 아래 행 → 헤더 (열 이름)
  - 헤더 바로 아래 3행 → 데이터 샘플 (AI 판단 근거)

출력 분기:
  - 시트 ≤ 10개: stdout 으로 전체 상세 출력 (A 방식)
  - 시트 ≥ 11개: 파일(./bunyang-output/scan.txt)에 상세 저장 + stdout 요약 (C 방식)
    (--out 인자로 경로 오버라이드 가능)
"""

import sys
import re
import argparse
from pathlib import Path

# [표] / [차트] / [차트:line] 프리픽스 파싱
BLOCK_PREFIX_RE = re.compile(r'^\[(표|차트)(?::([A-Za-z]+))?\]\s*')

try:
    import openpyxl
except ImportError:
    print("openpyxl 가 필요합니다: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def first_nonempty(row):
    for c in row:
        if c is not None and str(c).strip() != '':
            return str(c).strip()
    return ''


def trim_trailing(row):
    arr = list(row)
    while arr and (arr[-1] is None or str(arr[-1]).strip() == ''):
        arr.pop()
    return arr


def is_empty_row(row):
    if row is None:
        return True
    return all(c is None or str(c).strip() == '' for c in row)


def infer_type(cells):
    """열에 담긴 값들이 숫자형인지 텍스트형인지 추정"""
    if not cells:
        return 'empty'
    sample = [str(c).replace(',', '').replace('%', '').replace(' ', '')
              for c in cells if c is not None and str(c).strip()]
    if not sample:
        return 'empty'
    numeric = 0
    for s in sample:
        try:
            float(s)
            numeric += 1
        except ValueError:
            pass
    ratio = numeric / len(sample)
    if ratio >= 0.7:
        return '숫자형'
    if ratio <= 0.2:
        return '텍스트형'
    return '혼합형'


def scan_sheet(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return {'main': '', 'sub': '(없음)', 'blocks': [], 'needs_auto_message': False}

    # 1행 = main (위치 고정 — 빈 행이어도 skip 안 함)
    main = first_nonempty(rows[0]) if rows else ''
    # 2행 = sub (위치 고정)
    sub = first_nonempty(rows[1]) if len(rows) >= 2 else ''
    sub_display = sub if sub else '(없음)'
    needs_auto_message = (not main) and (not sub)

    # [표] / [차트] 행 찾기 + 헤더/샘플 추출
    blocks = []
    for r in range(len(rows)):
        first = first_nonempty(rows[r])
        m = BLOCK_PREFIX_RE.match(first)
        if not m:
            continue

        kind_ko = m.group(1)                    # '표' | '차트'
        chart_type = (m.group(2) or '').lower() # 'line' 등 (없으면 '')
        title_raw = first                       # "[차트:line] 매매가 추이"
        title = first[m.end():].strip()         # "매매가 추이"
        kind = 'chart' if kind_ko == '차트' else 'table'

        # 헤더 = 다음 non-empty 행
        header_row_idx = None
        for rr in range(r + 1, len(rows)):
            if not is_empty_row(rows[rr]):
                header_row_idx = rr
                break
        if header_row_idx is None:
            continue
        headers = [str(c).strip() if c is not None else ''
                   for c in trim_trailing(rows[header_row_idx])]

        # 샘플 = 헤더 다음 3행 (빈 행 만나면 중단)
        samples = []
        for rr in range(header_row_idx + 1, min(header_row_idx + 4, len(rows))):
            if is_empty_row(rows[rr]):
                break
            row_cells = trim_trailing(rows[rr])
            if not row_cells:
                break
            samples.append([str(c).strip() if c is not None else ''
                            for c in row_cells])

        # 데이터 컬럼 타입 (첫 컬럼 제외하고 나머지로 판단 — 첫 컬럼은 라벨인 경우가 많음)
        value_cells = []
        for s in samples:
            value_cells.extend(s[1:] if len(s) > 1 else s)
        col_type = infer_type(value_cells)

        blocks.append({
            'title_raw': title_raw,
            'title': title,
            'kind': kind,
            'chart_type': chart_type,
            'headers': headers,
            'n_cols': len(headers),
            'col_type': col_type,
            'samples': samples
        })

    return {'main': main, 'sub': sub_display, 'blocks': blocks,
            'needs_auto_message': needs_auto_message}


def format_sheet(name, data):
    lines = []
    lines.append(f"=== {name} ===")
    main_display = data['main'] if data['main'] else '(없음)'
    lines.append(f"main: {main_display}")
    lines.append(f"sub:  {data['sub']}")
    if data.get('needs_auto_message'):
        lines.append("  ⚠ 메인·서브 모두 비어있음 — Claude 가 블록 내용을 보고 "
                     "deck.cjs 에 mainMessage/subhead 를 직접 기입해야 함")
    blocks = data['blocks']
    n = len(blocks)
    n_tb = sum(1 for b in blocks if b['kind'] == 'table')
    n_ch = sum(1 for b in blocks if b['kind'] == 'chart')
    if n == 0:
        lines.append("blocks: (없음)")
    else:
        counts = []
        if n_tb: counts.append(f"표 {n_tb}")
        if n_ch: counts.append(f"차트 {n_ch}")
        lines.append(f"blocks ({n}개: {', '.join(counts)}):")
        for b in blocks:
            lines.append(f"  {b['title_raw']}")
            lines.append(
                f"    열: {' / '.join(b['headers'])}"
                f"   ({b['n_cols']}열, {b['col_type']})")
            if b['samples']:
                lines.append("    샘플 3행:")
                for s in b['samples']:
                    lines.append("      " + ' | '.join(s))
    return '\n'.join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('excel', help='엑셀 파일 경로')
    ap.add_argument('--out', default='./bunyang-output/scan.txt',
                    help='시트 11개↑일 때 상세 저장 경로 (기본: ./bunyang-output/scan.txt)')
    ap.add_argument('--force-file', action='store_true',
                    help='시트 개수와 관계없이 파일 저장 강제')
    args = ap.parse_args()

    excel_path = Path(args.excel)
    if not excel_path.exists():
        print(f"엑셀 파일 없음: {excel_path}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet_names = wb.sheetnames
    n_sheets = len(sheet_names)

    # 전체 스캔
    scans = []
    total_blocks = 0
    warnings = []
    for name in sheet_names:
        data = scan_sheet(wb[name])
        scans.append((name, data))
        nb = len(data['blocks'])
        total_blocks += nb
        if nb >= 5:
            warnings.append(f"  ⚠ 시트 '{name}' 에 블록 {nb}개 → 슬라이드 분할 필요")
        if data.get('needs_auto_message'):
            warnings.append(f"  ⚠ 시트 '{name}' 메인·서브 비어있음 → deck.cjs 에 직접 기입 필요")

    # 포맷
    full = '\n\n'.join(format_sheet(name, data) for name, data in scans)

    # 분기: 10개 이하 → stdout, 11개 이상 → 파일 + 요약 stdout
    use_file = (n_sheets >= 11) or args.force_file

    if use_file:
        out_path = Path(args.out)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(full, encoding='utf-8')
        print(f"📋 엑셀 스캔 완료 — {n_sheets}개 시트 / 총 {total_blocks}개 블록")
        print(f"   상세 결과: {out_path}")
        if warnings:
            print("\n경고:")
            for w in warnings:
                print(w)
    else:
        # 10개 이하 → 전체 stdout
        print(full)
        print(f"\n📋 합계: {n_sheets}개 시트 / {total_blocks}개 블록")
        if warnings:
            print("\n경고:")
            for w in warnings:
                print(w)


if __name__ == '__main__':
    main()
